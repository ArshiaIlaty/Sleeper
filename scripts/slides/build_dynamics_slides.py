"""Add two slides — sleep-stage dynamics and feature significance — to the
dataset section of the Scoring Metrics deck.

Reuses the styling helpers from build_dataset_slides.py so the new slides match
the Edwards template exactly. Inserts them right after "Impaired patients sleep
measurably worse" (deepening that finding with the transition mechanism and the
formal significance test). Idempotent: refuses to double-insert.

    python3 scripts/slides/build_dynamics_slides.py
"""
import os
import argparse

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

import build_dataset_slides as B   # same directory


def slide_dynamics(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[B.LAYOUT_CONTENT])
    B.set_title(s, "How impaired sleep breaks down, stage by stage")
    B.add_eyebrow(s, "DATASET · STAGE-TRANSITION DYNAMICS")
    B.fit_picture(s, os.path.join(figdir, "fig7_transition_matrix.png"),
                  0.55, 1.6, 7.1, 2.75, halign="center", valign="top")
    B.add_text(s, 0.6, 4.5, 7.1, 0.9,
               [[("Each cell = P(row stage → next epoch's stage). ", False, B.INK2, 12),
                 ("Impaired patients show a destabilized deep sleep", True, B.RED, 12),
                 (" — N3→N3 drops 5.4 pts while N3→N2 rises 6.0 — and less stable REM "
                  "(−4.8). The mechanism behind their fragmentation.", False, B.INK2, 12)]],
               base=12)
    B.add_footer(s)
    B.add_notes(s,
        "We can go deeper than 'how much' of each stage — into how the night "
        "moves between stages. This is the epoch-to-epoch transition matrix: each "
        "cell is the probability that a stage is followed by another. The left "
        "panel is the whole cohort — sleep is mostly stable, the diagonal is "
        "high. The right panel is what's different about the impaired patients, "
        "in percentage points. The striking cell is deep sleep: their N3 is far "
        "less stable — it drops back to N2 six points more often and holds only "
        "5 points less — and their REM is less stable too. So the fragmentation "
        "we saw in the summary metrics has a concrete mechanism: impaired "
        "patients can't hold onto their deep, restorative sleep.")
    return s


def slide_significance(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[B.LAYOUT_CONTENT])
    B.set_title(s, "Which features actually separate the groups?")
    B.add_eyebrow(s, "DATASET · FEATURE SIGNIFICANCE (χ² / t-test, FDR-corrected)")
    B.fit_picture(s, os.path.join(figdir, "fig8_feature_significance.png"),
                  0.5, 1.62, 3.85, 3.55, halign="left", valign="top")
    bullets = [
        [("14 of 48 features", True, B.INK, 12),
         (" are significant at FDR q<0.05.", False, B.INK2, 12)],
        [("Age dominates", True, B.RED, 12),
         (" (Cohen d = 1.08) — exactly why the metric discounts it.", False, B.INK2, 12)],
        [("Then: ", False, B.INK2, 12), ("limb movements, ↓REM, ↑WASO/Wake, "
          "↓efficiency & entropy, fewer REM periods, ↓N3.", True, B.INK, 12)],
        [("AHI is NOT significant", True, B.RED, 12),
         (" (d=0.19) — signal is in sleep architecture & continuity, not apnea load.",
          False, B.INK2, 12)],
        [("Method: ", False, B.INK2, 12), ("Welch t (numeric) + χ² (categorical), "
          "Benjamini–Hochberg FDR.", False, B.INK2, 11.5)],
    ]
    y = 1.72
    for b in bullets:
        B.add_text(s, 4.55, y, 3.1, 0.85, [b], base=12)
        y += 0.68
    B.add_footer(s)
    B.add_notes(s,
        "Finally, a formal answer to 'which features actually matter.' We tested "
        "every feature against the impairment label — a t-test for numeric "
        "features, chi-square for categorical — and corrected for testing 48 of "
        "them at once with a false-discovery-rate adjustment. Fourteen survive at "
        "q below 0.05. Age has by far the largest effect, a full standard "
        "deviation of separation — which is precisely the shortcut the "
        "age-conditioned metric strips away, so we can't lean on it. After age "
        "come the sleep features: periodic limb movements, reduced REM, more time "
        "awake, lower efficiency and stage entropy, fewer REM cycles, less deep "
        "sleep. The important negative result is that the apnea-hypopnea index is "
        "NOT significant — so the real signal lives in sleep architecture and "
        "continuity, not in how many breathing events a patient has. That "
        "directly justifies our feature design.")
    return s


def already_present(prs):
    for s in prs.slides:
        if s.shapes.title is not None and \
           (s.shapes.title.text or "").strip() == "Which features actually separate the groups?":
            return True
    return False


def _title_index(prs, title):
    for i, s in enumerate(prs.slides):
        if s.shapes.title is not None and (s.shapes.title.text or "").strip() == title:
            return i
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", default="PhysioNet Challenge 2026 - Scoring Metrics.pptx")
    ap.add_argument("--figures", default="paper/figures")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    if already_present(prs):
        raise SystemExit("Dynamics/significance slides already present — nothing to do.")

    anchor = _title_index(prs, "Impaired patients sleep measurably worse")
    if anchor is None:
        raise SystemExit("Could not find the anchor slide; run build_dataset_slides.py first.")

    slide_dynamics(prs, args.figures)
    slide_significance(prs, args.figures)
    B.reorder_after(prs, 2, anchor)
    prs.save(args.pptx)
    print(f"Added 2 slides after index {anchor}; deck now has {len(prs.slides)} slides.")


if __name__ == "__main__":
    main()
