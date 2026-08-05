"""Append a "Dataset behind the metric" section to the Scoring Metrics deck.

Adds one section-header slide plus six content slides (cohort, confounders,
biosignals/montages, sleep architecture, CI-vs-non-CI signal, data quality) to
`PhysioNet Challenge 2026 - Scoring Metrics.pptx`, wired with the same
Edwards-branded styling the deck already uses, plus speaker notes on every
slide. It also inserts the new section into the agenda and reorders it to land
right before the "What this means for our approach" takeaways.

All numbers come from eda/DATASET_REPORT.md (the verified EDA run); figures are
the vetted PNGs in paper/figures/. Idempotent: re-running detects the section
header and refuses to double-insert.

    python3 scripts/slides/build_dataset_slides.py \
        --pptx "PhysioNet Challenge 2026 - Scoring Metrics.pptx" \
        --figures paper/figures
"""
import os
import struct
import argparse

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Edwards palette (lifted verbatim from the existing deck) -----------------
RED       = RGBColor(0xC8, 0x10, 0x2E)   # primary accent / eyebrow
BLUE      = RGBColor(0x34, 0x65, 0x83)   # secondary accent
INK       = RGBColor(0x2A, 0x2E, 0x30)   # heading / strong body
INK2      = RGBColor(0x50, 0x57, 0x59)   # body gray
MUTED     = RGBColor(0x89, 0x8D, 0x8D)   # footer
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CARD      = RGBColor(0xF2, 0xF4, 0xF4)   # neutral card fill
CARD_BLUE = RGBColor(0xE9, 0xEE, 0xF0)   # accent card fill
HDR_GRAY  = RGBColor(0x50, 0x57, 0x59)   # table header
ROW_PINK  = RGBColor(0xFB, 0xE9, 0xEC)   # highlighted table row
ROW_ALT   = RGBColor(0xF7, 0xF8, 0xF8)

FONT = "Arial"
FOOTER_TXT = "PhysioNet Challenge 2026  ·  Scoring Metrics"

LAYOUT_CONTENT = 12   # "Title and Content"
LAYOUT_SECTION = 23   # "1_Section Header"


# ---- helpers ------------------------------------------------------------------
def png_aspect(path):
    """width/height of a PNG without a decoding dependency."""
    with open(path, "rb") as f:
        f.read(16)
        w, h = struct.unpack(">II", f.read(8))
    return w / h


def _set_runs(tf, spec, base_size=13, align=PP_ALIGN.LEFT):
    """Fill a text frame from a list of paragraphs; each paragraph is a list of
    (text, bold, color, size?) run tuples."""
    tf.word_wrap = True
    for pi, runs in enumerate(spec):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        for run in runs:
            text, bold, color = run[0], run[1], run[2]
            size = run[3] if len(run) > 3 else base_size
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def add_eyebrow(slide, text, color=RED):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.12), Inches(7.05), Inches(0.3))
    _set_runs(tb.text_frame, [[(text, True, color, 11)]])
    return tb


def add_footer(slide):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(5.28), Inches(6.6), Inches(0.25))
    _set_runs(tb.text_frame, [[(FOOTER_TXT, False, MUTED, 8)]])
    return tb


def add_text(slide, l, t, w, h, spec, base=13, align=PP_ALIGN.LEFT, anchor=None):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    if anchor is not None:
        tb.text_frame.vertical_anchor = anchor
    _set_runs(tb.text_frame, spec, base, align)
    return tb


def add_card(slide, l, t, w, h, fill=CARD, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    return shp


def add_stat_tile(slide, l, t, w, h, big, big_color, label):
    add_card(slide, l, t, w, h, CARD)
    tb = slide.shapes.add_textbox(Inches(l + 0.12), Inches(t + 0.1),
                                  Inches(w - 0.24), Inches(h - 0.2))
    _set_runs(tb.text_frame,
              [[(big, True, big_color, 21)],
               [(label, False, INK2, 9.5)]])
    return tb


def add_number_chip(slide, l, t, digit, size=0.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(l), Inches(t), Inches(size), Inches(0.55))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RED
    shp.line.fill.background()
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.25
    except Exception:
        pass
    tf = shp.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = digit
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return shp


def fit_picture(slide, path, box_l, box_t, box_w, box_h, halign="center", valign="middle"):
    """Place an image scaled to fit inside a box, preserving aspect."""
    a = png_aspect(path)
    w = box_w
    h = w / a
    if h > box_h:
        h = box_h
        w = h * a
    if halign == "center":
        l = box_l + (box_w - w) / 2
    elif halign == "right":
        l = box_l + (box_w - w)
    else:
        l = box_l
    if valign == "middle":
        t = box_t + (box_h - h) / 2
    elif valign == "bottom":
        t = box_t + (box_h - h)
    else:
        t = box_t
    return slide.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))


def style_cell(cell, text, bold, color, fill, size=10.5, align=PP_ALIGN.LEFT):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def add_table(slide, l, t, w, rows, cols, col_w, data, row_h=0.32, header_h=0.34):
    """data: list of rows; row 0 is the header. Each cell is (text, align)."""
    h = header_h + row_h * (rows - 1)
    gf = slide.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for ri in range(1, rows):
        tbl.rows[ri].height = Inches(row_h)
    for ri in range(rows):
        for ci in range(cols):
            text, align = data[ri][ci]
            if ri == 0:
                style_cell(tbl.cell(ri, ci), text, True, WHITE, HDR_GRAY, 10.5, align)
            else:
                fill = ROW_ALT if ri % 2 == 0 else WHITE
                style_cell(tbl.cell(ri, ci), text, False, INK, fill, 10.5, align)
    return gf


def set_title(slide, text):
    slide.shapes.title.text = text
    return slide.shapes.title


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---- slide builders -----------------------------------------------------------
def slide_section(prs):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_SECTION])
    set_title(s, "The dataset behind the metric")
    add_notes(s,
        "Now that we understand how we're scored, let's ground it in the actual "
        "data. Everything on the next few slides comes from a full statistical "
        "pass over all 1,103 recordings on the challenge dataset — demographics, "
        "the raw biosignals, and the CAISR sleep annotations. I want you to see "
        "two things: first, that the two problems the metric was built to handle "
        "— class imbalance and age confounding — are real and measurable in our "
        "cohort; and second, that there is genuine sleep signal separating "
        "impaired from unimpaired patients that a model can learn.")
    return s


def slide_cohort(prs):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, "1,103 recordings across three clinical sites")
    add_eyebrow(s, "DATASET · COHORT AT A GLANCE")
    # 2x2 stat tiles on the left
    tw, th, gap = 1.66, 0.95, 0.14
    x0, y0 = 0.6, 1.62
    add_stat_tile(s, x0, y0, tw, th, "1,103", INK, "patients · one overnight session each")
    add_stat_tile(s, x0 + tw + gap, y0, tw, th, "7.6%", RED, "impairment prevalence · 84 positive")
    add_stat_tile(s, x0, y0 + th + gap, tw, th, "62 ± 8.5", BLUE, "age (yr) · range 50–88")
    add_stat_tile(s, x0 + tw + gap, y0 + th + gap, tw, th, "3", INK, "clinical sites · BIDMC · Kaiser · Emory")
    # per-site table on the right
    tbl_l = 4.35
    data = [
        [("Site", PP_ALIGN.LEFT), ("Patients", PP_ALIGN.RIGHT), ("Prevalence", PP_ALIGN.RIGHT)],
        [("BIDMC", PP_ALIGN.LEFT), ("857", PP_ALIGN.RIGHT), ("6.5%", PP_ALIGN.RIGHT)],
        [("Kaiser", PP_ALIGN.LEFT), ("192", PP_ALIGN.RIGHT), ("10.4%", PP_ALIGN.RIGHT)],
        [("Emory", PP_ALIGN.LEFT), ("54", PP_ALIGN.RIGHT), ("14.8%", PP_ALIGN.RIGHT)],
        [("Cohort", PP_ALIGN.LEFT), ("1,103", PP_ALIGN.RIGHT), ("7.6%", PP_ALIGN.RIGHT)],
    ]
    add_table(s, tbl_l, 1.62, 3.3, 5, 3, [1.3, 1.0, 1.0], data)
    add_text(s, tbl_l, 3.7, 3.3, 0.6,
             [[("Prevalence rises 6.5% → 14.8% across sites — a source of "
                "distribution shift for leave-one-site-out validation.",
                False, INK2, 10.5)]], base=10.5)
    add_text(s, 0.6, 3.68, 3.5, 0.9,
             [[("Label", True, INK, 12), (" = ", False, INK2, 12),
               ("Cognitive_Impairment", True, INK, 12),
               ("  ·  84 positive / 1,019 negative, 0 missing.", False, INK2, 12)]],
             base=12)
    add_footer(s)
    add_notes(s,
        "Here is the cohort. 1,103 patients, one overnight recording each, drawn "
        "from three sites — Beth Israel in Boston, Kaiser, and Emory. The "
        "headline number for the metric is prevalence: only 7.6% of patients are "
        "cognitively impaired — 84 positives against a thousand-plus negatives. "
        "That is exactly the class imbalance the Reward metric is built to "
        "reward you for handling. Notice too that prevalence is not uniform "
        "across sites — from 6.5% at BIDMC up to nearly 15% at Emory — which is "
        "why we validate leave-one-site-out, not just by random fold. Age is "
        "tightly clustered around 62, with a long tail to the late 80s.")
    return s


def slide_confounders(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, "Both problems the metric fixes are visible in our data")
    add_eyebrow(s, "DATASET · WHY THE METRIC IS BUILT THIS WAY")
    fit_picture(s, os.path.join(figdir, "fig1_age_prevalence.png"),
                0.55, 1.65, 3.55, 3.35, halign="left")
    fit_picture(s, os.path.join(figdir, "fig2_site_prevalence.png"),
                4.15, 1.65, 3.5, 3.35, halign="left")
    add_text(s, 0.6, 4.95, 7.1, 0.5,
             [[("Prevalence climbs 2% → 36% across age bands", True, RED, 11.5),
               (" (→ age-conditioned AUROC) and shifts by site ", False, INK2, 11.5),
               ("6.5% → 14.8%", True, INK, 11.5),
               (" (→ leave-one-site-out validation).", False, INK2, 11.5)]],
             base=11.5)
    add_footer(s)
    add_notes(s,
        "This is the empirical case for the whole scoring design. On the left, "
        "prevalence of cognitive impairment by age band: 2% for patients in "
        "their fifties, rising to 36% for those 80 and older — an eighteen-fold "
        "swing. A lazy model could score well just by learning 'older equals "
        "higher risk,' which is precisely why the challenge conditions AUROC on "
        "age: it strips that shortcut away. On the right, prevalence by site: "
        "from 6.5% at BIDMC to almost 15% at Emory. The error bars are wide at "
        "the smaller sites, but the shift is real and it's why we don't trust a "
        "single random split. So the two things we said in the abstract — "
        "imbalance and age confounding — are not hypothetical; they are right "
        "here in the data.")
    return s


def slide_biosignals(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, "Rich signals, but wildly heterogeneous montages")
    add_eyebrow(s, "DATASET · BIOSIGNALS (header scan of 1,103 EDFs)")
    fit_picture(s, os.path.join(figdir, "fig5_montage_heterogeneity.png"),
                0.55, 1.6, 3.9, 3.5, halign="left")
    bullets = [
        [("EKG present in 98.9%", True, INK, 12), (" of recordings; core EEG "
          "(F3/C3/O1…) and SaO₂ in ~81%.", False, INK2, 12)],
        [("16 – 88 channels", True, INK, 12), (" per file; sampling rates span ",
          False, INK2, 12), ("20 – 512 Hz.", True, INK, 12)],
        [("BIDMC alone uses 42 distinct channel sets", True, RED, 12),
         ("; Kaiser 9, Emory 8.", False, INK2, 12)],
        [("Same sensor, many labels", True, INK, 12),
         (" (SaO2 vs SpO2, CHEST vs Thorax) and mixed units.", False, INK2, 12)],
        [("Implication:", True, RED, 12), (" features are computed per channel "
          "role and must degrade gracefully when a channel is absent.",
          False, INK2, 12)],
    ]
    y = 1.75
    for b in bullets:
        add_text(s, 4.65, y, 3.0, 0.8, [b], base=12)
        y += 0.66
    add_footer(s)
    add_notes(s,
        "The biosignals are rich but messy. A header-only scan of all 1,103 EDF "
        "files — we never load the 170 gigabytes of samples — shows EKG in "
        "almost every recording, and the core EEG derivations plus oxygen "
        "saturation in about four out of five. But the montages are wildly "
        "inconsistent: files range from 16 to 88 channels, sampling rates from "
        "20 to 512 hertz, and BIDMC alone uses 42 different channel sets. The "
        "same sensor shows up under different labels and units. The practical "
        "consequence, and a real design constraint for us, is that we extract "
        "features by channel role, not by exact label, and every feature has to "
        "degrade gracefully when its channel simply isn't there.")
    return s


def slide_sleep(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, "Fragmented sleep and heavy disordered breathing")
    add_eyebrow(s, "DATASET · SLEEP ARCHITECTURE (CAISR annotations)")
    fit_picture(s, os.path.join(figdir, "fig3_stage_composition.png"),
                0.55, 1.55, 7.1, 1.75, halign="center", valign="top")
    tw, th, gap = 1.66, 0.95, 0.14
    x0, y0 = 0.6, 3.5
    add_stat_tile(s, x0, y0, tw, th, "73.6%", BLUE, "sleep efficiency (median 76.8%)")
    add_stat_tile(s, x0 + (tw + gap), y0, tw, th, "47.9", RED, "AHI events / h of sleep")
    add_stat_tile(s, x0 + 2 * (tw + gap), y0, tw, th, "39.7", INK, "arousal index / h")
    add_stat_tile(s, x0 + 3 * (tw + gap), y0, tw, th, "332", INK, "total sleep time (min)")
    add_footer(s)
    add_notes(s,
        "These come from the CAISR automated sleep annotations, which we read "
        "for all 1,090 recordings that have them. The stage bar at the top is "
        "the cohort-average hypnogram: about a quarter of the night is wake, "
        "nearly half is N2, and only 9% is deep N3 and 12% REM. That is a "
        "fragmented, shallow sleep profile — consistent with a clinical sleep-lab "
        "population. The tiles below make it concrete: median sleep efficiency in "
        "the mid-70s, and a strikingly high apnea-hypopnea index of 48 events per "
        "hour of sleep, with an arousal index near 40. This is a heavily "
        "sleep-disordered cohort, which matters because sleep-disordered "
        "breathing is itself linked to cognitive risk. One caveat: CAISR agrees "
        "with the expert scorer on about 76% of epochs, so we treat these as "
        "strong but imperfect labels.")
    return s


def slide_signal(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, "Impaired patients sleep measurably worse")
    add_eyebrow(s, "DATASET · THE LEARNABLE SIGNAL (CI vs non-CI)")
    fit_picture(s, os.path.join(figdir, "fig4_ci_vs_noci.png"),
                0.55, 1.55, 7.1, 3.65, halign="center", valign="top")
    add_footer(s)
    add_notes(s,
        "This is the slide that says a model can actually work. Each panel "
        "compares the impaired group in red against the unimpaired group in "
        "blue. Impaired patients show lower sleep efficiency, shorter total "
        "sleep, and notably less deep N3 and REM sleep. The distributions "
        "overlap — this is not a clean separation, no single feature is a "
        "biomarker — but the shifts are consistent and in the direction sleep "
        "medicine would predict. Apnea burden and arousal index are more "
        "similar between the groups, which tells us the discriminative signal "
        "lives mainly in sleep architecture and continuity, not in the "
        "breathing indices. That is exactly why we lean on engineered "
        "architecture features rather than raw event counts.")
    return s


def slide_quality(prs, figdir):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, "What's present, and what we have to handle")
    add_eyebrow(s, "DATASET · COMPLETENESS & QUALITY")
    fit_picture(s, os.path.join(figdir, "fig6_data_completeness.png"),
                0.55, 1.6, 3.95, 3.5, halign="left")
    bullets = [
        [("Age & Sex complete", True, INK, 12), ("; race and ethnicity "
          "essentially complete.", False, INK2, 12)],
        [("BMI missing in 76%", True, RED, 12),
         (", Time_to_Event in 92% (recorded for positives only).",
          False, INK2, 12)],
        [("1,090 / 1,103 have CAISR", True, INK, 12),
         (" — 13 recordings fall back to global features only.",
          False, INK2, 12)],
        [("18 short recordings", True, INK, 12),
         (" (<4 h; 9 under 1 h) are flagged for the training pipeline.",
          False, INK2, 12)],
        [("Demographics ↔ EDF linkage is exact", True, INK, 12),
         (": 0 orphans in either direction.", False, INK2, 12)],
    ]
    y = 1.75
    for b in bullets:
        add_text(s, 4.7, y, 2.95, 0.8, [b], base=12)
        y += 0.66
    add_footer(s)
    add_notes(s,
        "Finally, data quality, because it drives real modeling decisions. The "
        "demographics we can count on — age and sex are complete, race and "
        "ethnicity nearly so. But BMI is missing for three-quarters of patients "
        "and time-to-event for over 90%, recorded only for the positives, so "
        "neither can be a load-bearing feature. On the signal side, 1,090 of "
        "1,103 recordings carry CAISR annotations; the 13 that don't fall back "
        "to global features. Eighteen recordings are too short — nine under an "
        "hour — and we flag those for the pipeline. The one thing that is "
        "perfectly clean is the linkage between the demographics table and the "
        "EDF files: zero orphans in either direction, so every patient's "
        "metadata lines up with their recording.")
    return s


# ---- agenda + ordering --------------------------------------------------------
def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def update_agenda(prs):
    """Insert 'The dataset behind the metric' as agenda item 6, renumbering the
    old item 6 ('What it means for us') to 7, with compressed spacing."""
    agenda = prs.slides[1]
    # (chip_name, text_name) for the six existing items, in reading order.
    pairs = [
        ("Rounded Rectangle 4", "TextBox 5"),
        ("Rounded Rectangle 6", "TextBox 7"),
        ("Rounded Rectangle 8", "TextBox 9"),
        ("Rounded Rectangle 10", "TextBox 11"),
        ("Rounded Rectangle 12", "TextBox 13"),
        ("Rounded Rectangle 14", "TextBox 15"),   # 'What it means for us'
    ]
    tops = [1.55, 2.07, 2.59, 3.11, 3.63, 4.15, 4.67]  # 7 slots, 0.52 spacing
    # slots 0..4 -> items 1..5 unchanged; slot 5 -> NEW item; slot 6 -> old item 6
    order = [0, 1, 2, 3, 4, 6]  # target slot for each existing pair
    for (chip_name, tb_name), slot in zip(pairs, order):
        chip = find_shape(agenda, chip_name)
        tb = find_shape(agenda, tb_name)
        if chip is not None:
            chip.top = Inches(tops[slot])
        if tb is not None:
            tb.top = Inches(tops[slot])
    # renumber the moved last item's chip 6 -> 7
    last_chip = find_shape(agenda, "Rounded Rectangle 14")
    if last_chip is not None:
        for p in last_chip.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() == "6":
                    r.text = "7"
    # add new item at slot 5
    add_number_chip(agenda, 0.6, tops[5], "6")
    add_text(agenda, 1.25, tops[5], 6.3, 0.55,
             [[("The dataset behind the metric  ", True, INK, 13),
               ("The cohort the model actually sees — and the confounders it must beat",
                False, INK2, 13)]],
             base=13)


def reorder_after(prs, n_new, after_index):
    """Move the last n_new slides to sit immediately after `after_index`."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new = ids[-n_new:]
    for e in new:
        lst.remove(e)
    ref = list(lst)[after_index + 1]  # element currently after the anchor
    for e in new:
        ref.addprevious(e)


def already_present(prs):
    for s in prs.slides:
        if s.shapes.title is not None and \
           (s.shapes.title.text or "").strip() == "The dataset behind the metric":
            return True
    return False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", default="PhysioNet Challenge 2026 - Scoring Metrics.pptx")
    ap.add_argument("--figures", default="paper/figures")
    ap.add_argument("--after-index", type=int, default=14,
                    help="insert the section right after this 0-based slide index")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    if already_present(prs):
        raise SystemExit("Dataset section already present — nothing to do.")

    figdir = args.figures
    builders = [
        lambda: slide_section(prs),
        lambda: slide_cohort(prs),
        lambda: slide_confounders(prs, figdir),
        lambda: slide_biosignals(prs, figdir),
        lambda: slide_sleep(prs, figdir),
        lambda: slide_signal(prs, figdir),
        lambda: slide_quality(prs, figdir),
    ]
    for b in builders:
        b()

    update_agenda(prs)
    reorder_after(prs, len(builders), args.after_index)
    prs.save(args.pptx)
    print(f"Added {len(builders)} slides; deck now has {len(prs.slides)} slides.")


if __name__ == "__main__":
    main()
